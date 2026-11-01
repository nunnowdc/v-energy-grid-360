"""Planilha Excel da fila de inspeção para a equipe de campo."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

# Só estas colunas vão para o arquivo: endereço sim, nome do titular não (LGPD)
COLUMNS = (
    ("Prioridade", "prioridade", 11),
    ("UC", "id_uc", 11),
    ("Score", "score", 8),
    ("Motivo da suspeita", "motivo_suspeita", 60),
    ("kWh recuperável/ano", "kwh_recuperavel_estimado_ano", 20),
    ("Município", "municipio", 24),
    ("Bairro", "bairro", 22),
    ("Logradouro", "logradouro", 32),
    ("Número", "numero", 9),
    ("Latitude", "latitude", 10),
    ("Longitude", "longitude", 10),
    ("Classe", "classe", 15),
    ("Ordem aberta", "ordem_aberta", 20),
)
KWH_FORMAT = "#,##0.00"


def build_queue_workbook(export: dict[str, object]) -> bytes:
    workbook = Workbook()
    queue = workbook.active
    queue.title = "Fila de inspeção"
    queue.append([title for title, _, _ in COLUMNS])
    # Célula a célula pela posição: queue[queue.max_row] varre a planilha inteira
    # a cada linha e deixava a fila real (~11,6 mil UCs) levando ~1 min
    for row, item in enumerate(export["itens"], start=2):
        for column, (_, key, _) in enumerate(COLUMNS, start=1):
            cell = queue.cell(row=row, column=column, value=item[key])
            # O arquivo sai do portal: texto da base que começa com "=" fica como
            # texto, nunca como fórmula que o Excel executaria ao abrir
            if cell.data_type == "f":
                cell.data_type = "s"

    for cell in queue[1]:
        cell.font = Font(bold=True)
    for index, (_, key, width) in enumerate(COLUMNS, start=1):
        letter = queue.cell(row=1, column=index).column_letter
        queue.column_dimensions[letter].width = width
        if key == "kwh_recuperavel_estimado_ano":
            for (cell,) in queue.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = KWH_FORMAT
    queue.freeze_panes = "A2"
    queue.auto_filter.ref = queue.dimensions

    filters = export["filtros"]
    about = workbook.create_sheet("Sobre")
    for row in (
        ["Data de referência", export["data_referencia"].strftime("%d/%m/%Y")],
        ["Município", filters["municipio"] or "Todos"],
        ["Bairro", filters["bairro"] or "Todos"],
        ["Classe", filters["classe"] or "Todas"],
        ["UCs", len(export["itens"])],
    ):
        about.append(row)
    about.column_dimensions["A"].width = 20
    about.column_dimensions["B"].width = 28

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
