"""Contrato HTTP da exportação da fila Gold em Excel para a equipe de campo."""

from contextlib import nullcontext
from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CABECALHO = [
    "Prioridade",
    "UC",
    "Score",
    "Motivo da suspeita",
    "kWh recuperável/ano",
    "Município",
    "Bairro",
    "Logradouro",
    "Número",
    "Latitude",
    "Longitude",
    "Classe",
    "Ordem aberta",
]


class ActiveOrders:
    def __init__(self, active=()):
        self.active = list(active)

    def get_active_orders(self, ids_uc):
        return [row for row in self.active if row["id_uc"] in ids_uc]


class ExportGoldFake:
    columns = (
        "prioridade",
        "id_uc",
        "score",
        "motivo_suspeita",
        "kwh_recuperavel_estimado_ano",
        "municipio",
        "bairro",
        "logradouro",
        "numero",
        "latitude",
        "longitude",
        "classe",
    )
    rows = (
        (
            9,
            "UC000009",
            60,
            "Motivo D",
            50.5,
            "Olinda",
            "Bairro Novo",
            "Rua do Sol",
            "12",
            -7.99,
            -34.84,
            "RESIDENCIAL",
        ),
        (
            2,
            "UC000002",
            90,
            "Motivo A; Motivo B",
            100.25,
            "Recife",
            "Boa Viagem",
            "Av. Boa Viagem",
            "1000",
            -8.12,
            -34.9,
            "RESIDENCIAL",
        ),
        (
            7,
            "UC000007",
            72,
            "Motivo C",
            20.0,
            "Recife",
            "Casa Forte",
            "Praça de Casa Forte",
            "S/N",
            -8.03,
            -34.92,
            "COMERCIAL",
        ),
    )

    def __init__(self, snapshot=date(2026, 6, 1), active=()):
        self.snapshot = snapshot
        self.orders = ActiveOrders(active)
        self.data = [dict(zip(self.columns, row, strict=True)) for row in self.rows]

    def get_catalog(self):
        return [
            ("Recife", "Boa Viagem"),
            ("Recife", "Casa Forte"),
            ("Olinda", "Bairro Novo"),
            ("Jaboatão", "Piedade"),
        ], ["RESIDENCIAL", "COMERCIAL"]

    def get_snapshot_date(self):
        return self.snapshot

    def get_export_items(self, municipio, bairro, classe):
        return sorted(
            (
                row
                for row in self.data
                if (municipio is None or row["municipio"] == municipio)
                and (bairro is None or row["bairro"] == bairro)
                and (classe is None or row["classe"] == classe)
            ),
            key=lambda row: (row["prioridade"], row["id_uc"]),
        )


def export_response(params=None, gold=None):
    from app.controllers.inspections import get_inspections_repository

    app.dependency_overrides[get_inspections_repository] = lambda: nullcontext(
        gold if gold is not None else ExportGoldFake()
    )
    try:
        return TestClient(app, raise_server_exceptions=False).get(
            "/api/inspections/queue/export", params=params
        )
    finally:
        app.dependency_overrides.clear()


def linhas(response, aba="Fila de inspeção"):
    planilha = load_workbook(BytesIO(response.content))[aba]
    return [list(linha) for linha in planilha.iter_rows(values_only=True)]


def test_baixa_xlsx_da_fila_inteira_na_prioridade_com_endereco():
    response = export_response()

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == (
        'attachment; filename="fila-inspecao_2026-06-01.xlsx"'
    )
    assert linhas(response) == [
        CABECALHO,
        [
            2,
            "UC000002",
            90,
            "Motivo A; Motivo B",
            100.25,
            "Recife",
            "Boa Viagem",
            "Av. Boa Viagem",
            "1000",
            -8.12,
            -34.9,
            "RESIDENCIAL",
            None,
        ],
        [
            7,
            "UC000007",
            72,
            "Motivo C",
            20,
            "Recife",
            "Casa Forte",
            "Praça de Casa Forte",
            "S/N",
            -8.03,
            -34.92,
            "COMERCIAL",
            None,
        ],
        [
            9,
            "UC000009",
            60,
            "Motivo D",
            50.5,
            "Olinda",
            "Bairro Novo",
            "Rua do Sol",
            "12",
            -7.99,
            -34.84,
            "RESIDENCIAL",
            None,
        ],
    ]


def test_nome_do_titular_nunca_vai_para_o_arquivo():
    gold = ExportGoldFake()
    for row in gold.data:
        row["nome_titular"] = "Fulana de Tal"

    response = export_response(gold=gold)

    assert response.status_code == 200
    celulas = [celula for linha in linhas(response) for celula in linha]
    assert "Fulana de Tal" not in celulas
    assert "nome_titular" not in celulas


def test_ordem_aberta_aparece_com_id_e_situacao():
    gold = ExportGoldFake(
        active=[
            {"id": 12, "id_uc": "UC000007", "status": "EM_CAMPO"},
            {"id": 3, "id_uc": "UC000009", "status": "ABERTA"},
        ]
    )

    response = export_response(gold=gold)

    assert [linha[-1] for linha in linhas(response)[1:]] == [
        None,
        "OI-00012 (Em campo)",
        "OI-00003 (Aberta)",
    ]


def test_respeita_os_filtros_e_registra_o_recorte_na_aba_sobre():
    response = export_response(
        {"municipio": "Recife", "bairro": "Boa Viagem", "classe": "RESIDENCIAL"}
    )

    assert response.status_code == 200
    assert [linha[1] for linha in linhas(response)[1:]] == ["UC000002"]
    assert linhas(response, "Sobre") == [
        ["Data de referência", "01/06/2026"],
        ["Município", "Recife"],
        ["Bairro", "Boa Viagem"],
        ["Classe", "RESIDENCIAL"],
        ["UCs", 1],
    ]


def test_sem_filtros_a_aba_sobre_diz_todos():
    response = export_response({"municipio": "", "bairro": "", "classe": ""})

    assert linhas(response, "Sobre") == [
        ["Data de referência", "01/06/2026"],
        ["Município", "Todos"],
        ["Bairro", "Todos"],
        ["Classe", "Todas"],
        ["UCs", 3],
    ]


def test_recorte_vazio_baixa_so_o_cabecalho():
    response = export_response({"municipio": "Jaboatão"})

    assert response.status_code == 200
    assert linhas(response) == [CABECALHO]


def test_filtros_invalidos_e_parametros_fora_do_contrato():
    cases = [
        {"bairro": "Boa Viagem"},
        {"municipio": "Olinda", "bairro": "Boa Viagem"},
        {"municipio": "Desconhecido"},
        {"classe": "DESCONHECIDA"},
        {"pagina": "1"},
        [("municipio", "Recife"), ("municipio", "Olinda")],
    ]
    for params in cases:
        assert export_response(params).status_code == 422, params


def test_snapshot_vazio_e_banco_fora_do_ar_dao_503_sem_vazar_detalhes():
    from sqlalchemy.exc import OperationalError

    response = export_response(gold=ExportGoldFake(snapshot=None))
    assert response.status_code == 503
    assert response.json() == {"detail": "Fila Gold indisponível"}

    class FailingGold(ExportGoldFake):
        def get_snapshot_date(self):
            raise OperationalError("SELECT segredo", {}, Exception("senha"))

    response = export_response(gold=FailingGold())
    assert response.status_code == 503
    assert response.json() == {"detail": "Gold indisponível"}
    assert "segredo" not in response.text


def test_texto_que_parece_formula_vai_como_texto_e_nao_executa():
    gold = ExportGoldFake()
    gold.data[0]["logradouro"] = '=HYPERLINK("http://example.invalid","Rua")'
    gold.data[1]["motivo_suspeita"] = "+1+1"

    response = export_response(gold=gold)

    planilha = load_workbook(BytesIO(response.content))["Fila de inspeção"]
    celulas = [celula for linha in planilha.iter_rows(min_row=2) for celula in linha]
    assert not [celula.coordinate for celula in celulas if celula.data_type == "f"]
    valores = [celula.value for celula in celulas]
    assert '=HYPERLINK("http://example.invalid","Rua")' in valores


def test_fila_do_tamanho_da_real_exporta_em_segundos():
    # A Gold tem ~11,6 mil UCs na fila; montar a planilha não pode crescer
    # de forma quadrática (antes levava ~1 min e o botão ficava em "Exportando…")
    from time import perf_counter

    gold = ExportGoldFake()
    modelo = gold.data[0]
    gold.data = [
        {**modelo, "id_uc": f"UC{indice:06d}", "prioridade": indice}
        for indice in range(1, 12_001)
    ]

    inicio = perf_counter()
    response = export_response(gold=gold)
    duracao = perf_counter() - inicio

    assert response.status_code == 200
    assert len(linhas(response)) == 12_001
    assert duracao < 15
