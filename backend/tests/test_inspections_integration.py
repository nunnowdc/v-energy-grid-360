"""Comparação somente leitura da fila HTTP com a Gold local."""

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.core.config import get_settings
from app.core.database import get_connection
from app.main import app


@pytest.mark.skipif(
    not get_settings().database_url, reason="DATABASE_URL não configurada"
)
def test_fila_http_preserva_snapshot_gold_e_resumo_do_recorte() -> None:
    with get_connection() as connection:
        total, energy, reference = connection.execute(
            text(
                "SELECT COUNT(*), SUM(kwh_recuperavel_estimado_ano), "
                "MAX(data_referencia) FROM fila_inspecao"
            )
        ).one()
        gold_rows = (
            connection.execute(
                text(
                    "SELECT prioridade, id_uc, score, motivo_suspeita, "
                    "kwh_recuperavel_estimado_ano, municipio, bairro, classe "
                    "FROM fila_inspecao ORDER BY prioridade, id_uc LIMIT 20"
                )
            )
            .mappings()
            .all()
        )

    response = TestClient(app).get("/api/inspections/queue")
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == total == body["resumo"]["ucs"]
    assert body["resumo"]["kwh_recuperavel_estimado_ano"] == pytest.approx(energy)
    assert body["data_referencia"] == reference.isoformat()
    assert body["pagina"] == 1
    assert body["limite"] == 20
    assert len(body["itens"]) == len(gold_rows)
    for actual, expected in zip(body["itens"], gold_rows, strict=True):
        for key in expected:
            if key == "kwh_recuperavel_estimado_ano":
                assert actual[key] == pytest.approx(expected[key])
            else:
                assert actual[key] == expected[key]
        if actual["ordem_aberta"] is not None:
            assert set(actual["ordem_aberta"]) == {"id_ordem", "status"}
            assert actual["ordem_aberta"]["status"] in {"ABERTA", "EM_CAMPO"}
        assert set(actual) == set(expected) | {"ordem_aberta"}


@pytest.mark.skipif(
    not get_settings().database_url, reason="DATABASE_URL não configurada"
)
def test_fila_filtrada_e_segunda_pagina_conciliam_com_gold() -> None:
    with get_connection() as connection:
        municipality = connection.execute(
            text(
                "SELECT municipio FROM fila_inspecao GROUP BY municipio ORDER BY COUNT(*) DESC LIMIT 1"
            )
        ).scalar_one()
        total, energy = connection.execute(
            text(
                "SELECT COUNT(*), SUM(kwh_recuperavel_estimado_ano) "
                "FROM fila_inspecao WHERE municipio = :municipio"
            ),
            {"municipio": municipality},
        ).one()
        second_page = connection.execute(
            text(
                "SELECT id_uc, prioridade FROM fila_inspecao "
                "WHERE municipio = :municipio ORDER BY prioridade, id_uc "
                "LIMIT 20 OFFSET 20"
            ),
            {"municipio": municipality},
        ).all()

    response = TestClient(app).get(
        "/api/inspections/queue", params={"municipio": municipality, "pagina": 2}
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == total
    assert body["resumo"]["kwh_recuperavel_estimado_ano"] == pytest.approx(energy)
    assert [(item["id_uc"], item["prioridade"]) for item in body["itens"]] == [
        (row.id_uc, row.prioridade) for row in second_page
    ]


@pytest.mark.skipif(
    not get_settings().database_url, reason="DATABASE_URL não configurada"
)
def test_filtros_combinados_consultam_colunas_da_fila_sem_duplicacao() -> None:
    with get_connection() as connection:
        sample = connection.execute(
            text(
                "SELECT municipio, bairro, classe FROM fila_inspecao "
                "WHERE municipio IS NOT NULL AND bairro IS NOT NULL "
                "AND classe IS NOT NULL LIMIT 1"
            )
        ).one()
        params = dict(sample._mapping)
        total, energy = connection.execute(
            text(
                "SELECT COUNT(*), SUM(kwh_recuperavel_estimado_ano) "
                "FROM fila_inspecao WHERE municipio = :municipio "
                "AND bairro = :bairro AND classe = :classe"
            ),
            params,
        ).one()
        first = connection.execute(
            text(
                "SELECT prioridade, id_uc FROM fila_inspecao "
                "WHERE municipio = :municipio AND bairro = :bairro "
                "AND classe = :classe ORDER BY prioridade, id_uc LIMIT 1"
            ),
            params,
        ).one()

    response = TestClient(app).get("/api/inspections/queue", params=params)
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == total
    assert body["resumo"]["ucs"] == total
    assert body["resumo"]["kwh_recuperavel_estimado_ano"] == pytest.approx(energy)
    assert (body["itens"][0]["prioridade"], body["itens"][0]["id_uc"]) == (
        first.prioridade,
        first.id_uc,
    )


@pytest.mark.skipif(
    not get_settings().database_url, reason="DATABASE_URL não configurada"
)
def test_exportacao_traz_a_fila_inteira_do_recorte_com_endereco() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    with get_connection() as connection:
        municipality = connection.execute(
            text(
                "SELECT municipio FROM fila_inspecao GROUP BY municipio ORDER BY COUNT(*) DESC LIMIT 1"
            )
        ).scalar_one()
        gold_rows = connection.execute(
            text(
                "SELECT id_uc, logradouro, numero FROM fila_inspecao "
                "WHERE municipio = :municipio ORDER BY prioridade, id_uc"
            ),
            {"municipio": municipality},
        ).all()

    response = TestClient(app).get(
        "/api/inspections/queue/export", params={"municipio": municipality}
    )
    assert response.status_code == 200
    planilha = load_workbook(BytesIO(response.content))["Fila de inspeção"]
    linhas = list(planilha.iter_rows(min_row=2, values_only=True))
    assert [(linha[1], linha[7], linha[8]) for linha in linhas] == [
        (row.id_uc, row.logradouro, row.numero) for row in gold_rows
    ]
